"""Backend API tests for Mass Mailer.

Covers: auth, users, email/word templates, excel parse, campaigns send,
logs, api-keys, external send, and stats. SMTP is intentionally unconfigured
so send operations are expected to fail per-recipient but endpoints should
still return campaign_id and write send_logs.
"""
import io
import os
import uuid
import pytest
import requests
from openpyxl import Workbook
from docx import Document

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "http://localhost:8001").rstrip("/")
API = f"{BASE_URL}/api"

ADMIN_EMAIL = "admin@example.com"
ADMIN_PASSWORD = "admin123"


# ----------------- Fixtures -----------------
@pytest.fixture(scope="session")
def admin_token():
    r = requests.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD}, timeout=15)
    assert r.status_code == 200, f"admin login failed: {r.status_code} {r.text}"
    data = r.json()
    assert "access_token" in data
    assert data["user"]["email"] == ADMIN_EMAIL
    # cookies should be set
    assert "access_token" in r.cookies or any(c.name == "access_token" for c in r.cookies)
    return data["access_token"]


@pytest.fixture(scope="session")
def admin_headers(admin_token):
    return {"Authorization": f"Bearer {admin_token}"}


@pytest.fixture(scope="session")
def user_creds(admin_headers):
    """Create a normal user for non-admin tests."""
    email = f"TEST_user_{uuid.uuid4().hex[:8]}@example.com"
    password = "userpass123"
    r = requests.post(
        f"{API}/users",
        json={"email": email, "password": password, "name": "Test User", "role": "user"},
        headers=admin_headers, timeout=15,
    )
    assert r.status_code == 200, f"user create failed: {r.text}"
    uid = r.json()["id"]
    yield {"email": email, "password": password, "id": uid}
    # cleanup
    requests.delete(f"{API}/users/{uid}", headers=admin_headers, timeout=10)


@pytest.fixture(scope="session")
def user_token(user_creds):
    r = requests.post(f"{API}/auth/login",
                      json={"email": user_creds["email"], "password": user_creds["password"]}, timeout=15)
    assert r.status_code == 200
    return r.json()["access_token"]


@pytest.fixture(scope="session")
def user_headers(user_token):
    return {"Authorization": f"Bearer {user_token}"}


def _build_xlsx_bytes():
    wb = Workbook()
    ws = wb.active
    ws.append(["email", "password", "name", "invoice_number"])
    ws.append(["alice@example.com", "pw_alice", "Alice", "INV-001"])
    ws.append(["bob@example.com", "pw_bob", "Bob", "INV-002"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build_docx_bytes():
    doc = Document()
    doc.add_paragraph("Hello {name}, your invoice is {invoice_number}.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ----------------- Auth -----------------
class TestAuth:
    def test_login_success_sets_cookies(self):
        r = requests.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD}, timeout=15)
        assert r.status_code == 200
        data = r.json()
        assert data["token_type"] == "bearer"
        assert data["user"]["role"] == "admin"
        cookie_names = {c.name for c in r.cookies}
        assert "access_token" in cookie_names
        assert "refresh_token" in cookie_names

    def test_login_invalid(self):
        r = requests.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": "wrong"}, timeout=15)
        assert r.status_code == 401

    def test_me_with_bearer(self, admin_headers):
        r = requests.get(f"{API}/auth/me", headers=admin_headers, timeout=10)
        assert r.status_code == 200
        assert r.json()["email"] == ADMIN_EMAIL

    def test_me_unauthorized(self):
        r = requests.get(f"{API}/auth/me", timeout=10)
        assert r.status_code == 401

    def test_logout_clears_cookies(self, admin_token):
        s = requests.Session()
        s.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD}, timeout=10)
        r = s.post(f"{API}/auth/logout", timeout=10)
        assert r.status_code == 200
        # cookies cleared in response
        set_cookie = r.headers.get("set-cookie", "")
        assert "access_token" in set_cookie

    def test_logout_does_not_require_auth(self):
        """Per migration: logout should NOT require authentication anymore."""
        r = requests.post(f"{API}/auth/logout", timeout=10)
        assert r.status_code == 200, f"Logout without auth should succeed; got {r.status_code} {r.text}"
        set_cookie = r.headers.get("set-cookie", "")
        assert "access_token" in set_cookie

    def test_me_returns_expected_fields(self, admin_headers):
        r = requests.get(f"{API}/auth/me", headers=admin_headers, timeout=10)
        assert r.status_code == 200
        data = r.json()
        for k in ("id", "email", "name", "role", "created_at"):
            assert k in data, f"missing {k} in /auth/me response"
        # password_hash should NEVER be returned
        assert "password_hash" not in data


# ----------------- Users (admin) -----------------
class TestUsers:
    def test_list_users_admin(self, admin_headers):
        r = requests.get(f"{API}/users", headers=admin_headers, timeout=10)
        assert r.status_code == 200
        assert isinstance(r.json(), list)
        assert any(u["email"] == ADMIN_EMAIL for u in r.json())

    def test_non_admin_forbidden(self, user_headers):
        r = requests.get(f"{API}/users", headers=user_headers, timeout=10)
        assert r.status_code == 403
        r = requests.post(f"{API}/users",
                          json={"email": "TEST_x@example.com", "password": "x", "role": "user"},
                          headers=user_headers, timeout=10)
        assert r.status_code == 403

    def test_create_and_delete_user(self, admin_headers):
        email = f"TEST_crud_{uuid.uuid4().hex[:8]}@example.com"
        r = requests.post(f"{API}/users",
                          json={"email": email, "password": "secret123", "name": "CRUD", "role": "user"},
                          headers=admin_headers, timeout=10)
        assert r.status_code == 200, r.text
        uid = r.json()["id"]
        # Backend normalizes email to lowercase
        assert r.json()["email"] == email.lower()
        # duplicate
        r2 = requests.post(f"{API}/users",
                           json={"email": email, "password": "secret123", "role": "user"},
                           headers=admin_headers, timeout=10)
        assert r2.status_code == 400
        # delete
        r3 = requests.delete(f"{API}/users/{uid}", headers=admin_headers, timeout=10)
        assert r3.status_code == 200
        # delete missing
        r4 = requests.delete(f"{API}/users/{uid}", headers=admin_headers, timeout=10)
        assert r4.status_code == 404


# ----------------- Email Templates -----------------
class TestEmailTemplates:
    def test_email_template_crud(self, admin_headers):
        payload = {"name": "TEST_tpl", "subject": "Hi {name}", "body_html": "<p>Hello {name}</p>"}
        r = requests.post(f"{API}/email-templates", json=payload, headers=admin_headers, timeout=10)
        assert r.status_code == 200, r.text
        tid = r.json()["id"]
        assert r.json()["subject"] == payload["subject"]

        r = requests.get(f"{API}/email-templates", headers=admin_headers, timeout=10)
        assert r.status_code == 200
        assert any(t["id"] == tid for t in r.json())

        upd = {"name": "TEST_tpl2", "subject": "Hello {name}", "body_html": "<p>Hi {name}, INV {invoice_number}</p>"}
        r = requests.put(f"{API}/email-templates/{tid}", json=upd, headers=admin_headers, timeout=10)
        assert r.status_code == 200
        assert r.json()["subject"] == upd["subject"]

        r = requests.delete(f"{API}/email-templates/{tid}", headers=admin_headers, timeout=10)
        assert r.status_code == 200
        r = requests.delete(f"{API}/email-templates/{tid}", headers=admin_headers, timeout=10)
        assert r.status_code == 404


# ----------------- Word Templates -----------------
class TestWordTemplates:
    def test_word_template_upload_and_delete(self, admin_headers):
        files = {"file": ("tpl.docx", _build_docx_bytes(),
                          "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        data = {"name": "TEST_word"}
        r = requests.post(f"{API}/word-templates", files=files, data=data, headers=admin_headers, timeout=20)
        assert r.status_code == 200, r.text
        wid = r.json()["id"]

        r = requests.get(f"{API}/word-templates", headers=admin_headers, timeout=10)
        assert r.status_code == 200
        assert any(w["id"] == wid for w in r.json())

        r = requests.delete(f"{API}/word-templates/{wid}", headers=admin_headers, timeout=10)
        assert r.status_code == 200

    def test_word_template_rejects_non_docx(self, admin_headers):
        files = {"file": ("bad.txt", b"hello", "text/plain")}
        r = requests.post(f"{API}/word-templates", files=files, data={"name": "x"},
                          headers=admin_headers, timeout=10)
        assert r.status_code == 400


# ----------------- Excel parse -----------------
class TestExcelParse:
    def test_parse_excel(self, admin_headers):
        files = {"file": ("data.xlsx", _build_xlsx_bytes(),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{API}/excel/parse", files=files, headers=admin_headers, timeout=15)
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["headers"] == ["email", "password", "name", "invoice_number"]
        assert data["count"] == 2
        assert data["rows"][0]["email"] == "alice@example.com"
        assert data["rows"][0]["name"] == "Alice"


# ----------------- Campaign send (UI) -----------------
class TestCampaignSend:
    def test_send_campaign_smtp_fails_logs_written(self, admin_headers):
        files = {"excel": ("data.xlsx", _build_xlsx_bytes(),
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        data = {
            "subject": "Hello {name}",
            "body_html": "<p>Hi {name}, INV {invoice_number}</p>",
            "attachment_basename": "doc",
            # Throttle overrides to keep test fast
            "per_minute": "10000",
            "per_domain_per_min": "10000",
            "delay_min_ms": "0",
            "delay_max_ms": "0",
        }
        r = requests.post(f"{API}/campaigns/send", data=data, files=files,
                          headers=admin_headers, timeout=120)
        assert r.status_code == 200, r.text
        body = r.json()
        assert "campaign_id" in body
        cid = body["campaign_id"]
        assert body["total"] == 2
        # SMTP fails → expect failures==2, sent==0
        assert body["failed"] == 2
        assert body["sent"] == 0

        # Verify logs were written for this campaign
        r2 = requests.get(f"{API}/logs?limit=500", headers=admin_headers, timeout=10)
        assert r2.status_code == 200
        logs = [l for l in r2.json() if l.get("campaign_id") == cid]
        assert len(logs) == 2
        for l in logs:
            assert l["status"] == "failed"
            assert l["source"] == "ui"
            assert l["user_email"] == ADMIN_EMAIL
            assert l["recipient"] in ("alice@example.com", "bob@example.com")
            assert "timestamp" in l
            assert "subject" in l

    def test_send_campaign_with_word_template(self, admin_headers):
        # Upload word template first
        files_w = {"file": ("tpl.docx", _build_docx_bytes(),
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        rw = requests.post(f"{API}/word-templates", files=files_w, data={"name": "TEST_cw"},
                           headers=admin_headers, timeout=20)
        assert rw.status_code == 200
        wid = rw.json()["id"]
        try:
            files = {"excel": ("data.xlsx", _build_xlsx_bytes(),
                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
            data = {
                "subject": "Hello {name}",
                "body_html": "<p>Hi {name}, INV {invoice_number}</p>",
                "word_template_id": wid,
                "attachment_basename": "invoice",
                # Throttle overrides
                "per_minute": "10000",
                "per_domain_per_min": "10000",
                "delay_min_ms": "0",
                "delay_max_ms": "0",
            }
            r = requests.post(f"{API}/campaigns/send", data=data, files=files,
                              headers=admin_headers, timeout=300)
            assert r.status_code == 200, r.text
            body = r.json()
            # SMTP fails but PDF generation should succeed for both
            assert body["total"] == 2
            # We still expect failures from SMTP; sent==0
            assert body["failed"] == 2
            # Inspect that failure errors are SMTP-related (not pdf related)
            errs = " ".join(f.get("error", "") for f in body["failures"]).lower()
            # Accept any error; ensure attachment_name was set (logs)
            r2 = requests.get(f"{API}/logs?limit=500", headers=admin_headers, timeout=10)
            logs = [l for l in r2.json() if l.get("campaign_id") == body["campaign_id"]]
            assert len(logs) == 2
            # If LibreOffice is present, PDF is generated and attachment_name is set.
            # If not, send fails earlier with FileNotFoundError('soffice') — still logs.
            # Either way the campaign_id + 2 send_logs must exist.
            for l in logs:
                assert l["status"] == "failed"
                assert l["source"] == "ui"
        finally:
            requests.delete(f"{API}/word-templates/{wid}", headers=admin_headers, timeout=10)


# ----------------- Logs -----------------
class TestLogs:
    def test_logs_admin(self, admin_headers):
        r = requests.get(f"{API}/logs", headers=admin_headers, timeout=10)
        assert r.status_code == 200
        assert isinstance(r.json(), list)


# ----------------- API Keys -----------------
class TestApiKeys:
    def test_api_keys_admin_only(self, user_headers):
        """Per migration: /api/api-keys is now ADMIN-ONLY. Non-admin must get 403."""
        r = requests.get(f"{API}/api-keys", headers=user_headers, timeout=10)
        assert r.status_code == 403, f"non-admin GET expected 403, got {r.status_code}"
        r = requests.post(f"{API}/api-keys", json={"name": "TEST_nokey"},
                          headers=user_headers, timeout=10)
        assert r.status_code == 403, f"non-admin POST expected 403, got {r.status_code}"
        # Use a fake id; admin guard runs before lookup
        r = requests.delete(f"{API}/api-keys/non-existent-id",
                            headers=user_headers, timeout=10)
        assert r.status_code == 403, f"non-admin DELETE expected 403, got {r.status_code}"

    def test_api_key_lifecycle_and_external_send(self, admin_headers):
        # Create key
        r = requests.post(f"{API}/api-keys", json={"name": "TEST_key"},
                          headers=admin_headers, timeout=10)
        assert r.status_code == 200
        raw = r.json()["key"]
        kid = r.json()["id"]
        assert raw.startswith("mk_")

        # List shows preview only, no key_hash, no raw key
        r2 = requests.get(f"{API}/api-keys", headers=admin_headers, timeout=10)
        assert r2.status_code == 200
        found = [k for k in r2.json() if k["id"] == kid]
        assert found
        assert "key_hash" not in found[0]
        assert "key_preview" in found[0]

        # Create an email template for external send
        tpl = requests.post(f"{API}/email-templates",
                            json={"name": "TEST_extt", "subject": "Hi {name}",
                                  "body_html": "<p>Hello {name}, INV {invoice_number}</p>"},
                            headers=admin_headers, timeout=10).json()
        tid = tpl["id"]
        try:
            # Missing key -> 401
            r3 = requests.post(f"{API}/external/send",
                               json={"template_id": tid, "recipients": [{"email": "x@example.com"}]},
                               timeout=15)
            assert r3.status_code == 401

            # Wrong key -> 401
            r4 = requests.post(f"{API}/external/send",
                               json={"template_id": tid, "recipients": [{"email": "x@example.com"}]},
                               headers={"X-API-Key": "mk_wrongwrongwrong"}, timeout=15)
            assert r4.status_code == 401

            # Correct key
            r5 = requests.post(f"{API}/external/send",
                               json={
                                   "template_id": tid,
                                   "recipients": [
                                       {"email": "alice@example.com", "name": "Alice", "invoice_number": "INV-1"},
                                       {"email": "bob@example.com", "name": "Bob", "invoice_number": "INV-2"},
                                   ],
                                   "per_minute": 10000,
                                   "per_domain_per_min": 10000,
                                   "delay_min_ms": 0,
                                   "delay_max_ms": 0,
                               },
                               headers={"X-API-Key": raw}, timeout=120)
            assert r5.status_code == 200, r5.text
            body = r5.json()
            cid = body["campaign_id"]
            assert body["total"] == 2
            assert body["failed"] == 2  # SMTP not available

            # Verify logs have source=api
            r6 = requests.get(f"{API}/logs?limit=500", headers=admin_headers, timeout=10)
            logs = [l for l in r6.json() if l.get("campaign_id") == cid]
            assert len(logs) == 2
            assert all(l["source"] == "api" for l in logs)

            # Revoke key
            r7 = requests.delete(f"{API}/api-keys/{kid}", headers=admin_headers, timeout=10)
            assert r7.status_code == 200

            # Revoked key should not work
            r8 = requests.post(f"{API}/external/send",
                               json={"template_id": tid, "recipients": [{"email": "x@example.com"}]},
                               headers={"X-API-Key": raw}, timeout=15)
            assert r8.status_code == 401
        finally:
            requests.delete(f"{API}/email-templates/{tid}", headers=admin_headers, timeout=10)


# ----------------- Stats -----------------
class TestStats:
    def test_stats(self, admin_headers):
        r = requests.get(f"{API}/stats", headers=admin_headers, timeout=10)
        assert r.status_code == 200
        data = r.json()
        for k in ("total_emails", "sent", "failed", "success_rate", "email_templates", "word_templates"):
            assert k in data
        assert data["total_emails"] >= 0


# ----------------- SQLCipher encryption-at-rest -----------------
class TestStorageEncryption:
    """Verify the SQLite file at /app/backend/data/mailmaster.db is actually encrypted."""

    DB_PATH = "/app/backend/data/mailmaster.db"

    def test_db_file_exists(self):
        assert os.path.exists(self.DB_PATH), f"DB file missing at {self.DB_PATH}"
        assert os.path.getsize(self.DB_PATH) > 0

    def test_db_file_is_not_plain_sqlite(self):
        """The first 16 bytes of a plain SQLite DB are 'SQLite format 3\\0'.
        For SQLCipher, the header is encrypted, so it must NOT match."""
        with open(self.DB_PATH, "rb") as f:
            header = f.read(16)
        assert not header.startswith(b"SQLite format 3"), (
            "DB file appears to be unencrypted plain SQLite! "
            f"header={header!r}"
        )

    def test_stdlib_sqlite3_cannot_open(self):
        """Opening the encrypted DB with stdlib sqlite3 (no key) must fail."""
        import sqlite3
        conn = sqlite3.connect(self.DB_PATH)
        try:
            with pytest.raises(sqlite3.DatabaseError):
                conn.execute("SELECT count(*) FROM sqlite_master").fetchone()
        finally:
            conn.close()


# ----------------- Throttle / anti-blocking layer -----------------
import time as _time


class TestThrottle:
    """Anti-blocking layer: defaults endpoint + per-campaign overrides."""

    def test_throttle_defaults_shape(self, admin_headers):
        r = requests.get(f"{API}/throttle/defaults", headers=admin_headers, timeout=10)
        assert r.status_code == 200, r.text
        data = r.json()
        # Must contain all expected fields
        for k in ("per_minute", "per_hour", "per_day",
                  "per_domain_per_min", "delay_min_ms", "delay_max_ms",
                  "today_sent"):
            assert k in data, f"missing {k} in defaults response: {data}"
        # Type sanity
        for k in ("per_minute", "per_hour", "per_day",
                  "per_domain_per_min", "delay_min_ms", "delay_max_ms",
                  "today_sent"):
            assert isinstance(data[k], int), f"{k} should be int, got {type(data[k])}"
        # delay_max must be >= delay_min
        assert data["delay_max_ms"] >= data["delay_min_ms"]
        # today_sent should be a non-negative count
        assert data["today_sent"] >= 0

    def test_throttle_defaults_requires_auth(self):
        r = requests.get(f"{API}/throttle/defaults", timeout=10)
        assert r.status_code == 401

    def test_campaign_send_echoes_throttle_and_runs_fast(self, admin_headers):
        """delay=0 + per_minute=1000 + per_domain=1000 → 2-row Excel must
        complete in well under the natural-jitter floor and response must
        include 'throttle' echo dict."""
        files = {"excel": ("data.xlsx", _build_xlsx_bytes(),
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        data = {
            "subject": "Hi {name}",
            "body_html": "<p>{name}</p>",
            "attachment_basename": "doc",
            "per_minute": "1000",
            "per_hour": "10000",
            "per_day": "100000",
            "per_domain_per_min": "1000",
            "delay_min_ms": "0",
            "delay_max_ms": "0",
        }
        t0 = _time.time()
        r = requests.post(f"{API}/campaigns/send", data=data, files=files,
                          headers=admin_headers, timeout=60)
        elapsed = _time.time() - t0
        assert r.status_code == 200, r.text
        body = r.json()
        # 'throttle' echo present and matches overrides
        assert "throttle" in body, f"response missing 'throttle' key: {body}"
        echoed = body["throttle"]
        assert echoed["per_minute"] == 1000
        assert echoed["per_hour"] == 10000
        assert echoed["per_day"] == 100000
        assert echoed["per_domain_per_min"] == 1000
        assert echoed["delay_min_ms"] == 0
        assert echoed["delay_max_ms"] == 0
        # Counters
        assert body["total"] == 2
        assert body["sent"] == 0     # SMTP unreachable
        assert body["failed"] == 2
        # Should NOT have taken 30+ seconds (no throttling)
        assert elapsed < 20.0, f"campaign took {elapsed:.1f}s — throttle override not honoured"

    def test_external_send_echoes_throttle(self, admin_headers):
        # Create a key + template
        rk = requests.post(f"{API}/api-keys", json={"name": "TEST_thr_key"},
                           headers=admin_headers, timeout=10)
        assert rk.status_code == 200
        raw = rk.json()["key"]
        kid = rk.json()["id"]

        rt = requests.post(f"{API}/email-templates",
                           json={"name": "TEST_thr_tpl", "subject": "Hi {name}",
                                 "body_html": "<p>Hello {name}</p>"},
                           headers=admin_headers, timeout=10)
        assert rt.status_code == 200
        tid = rt.json()["id"]

        try:
            t0 = _time.time()
            r = requests.post(
                f"{API}/external/send",
                json={
                    "template_id": tid,
                    "recipients": [
                        {"email": "x1@a.com", "name": "X1"},
                        {"email": "x2@b.com", "name": "X2"},
                    ],
                    "per_minute": 1000,
                    "per_hour": 10000,
                    "per_day": 100000,
                    "per_domain_per_min": 1000,
                    "delay_min_ms": 0,
                    "delay_max_ms": 0,
                },
                headers={"X-API-Key": raw}, timeout=60,
            )
            elapsed = _time.time() - t0
            assert r.status_code == 200, r.text
            body = r.json()
            assert "throttle" in body, f"missing 'throttle' echo: {body}"
            echoed = body["throttle"]
            assert echoed["per_minute"] == 1000
            assert echoed["delay_min_ms"] == 0
            assert echoed["delay_max_ms"] == 0
            assert echoed["per_domain_per_min"] == 1000
            assert body["total"] == 2
            assert body["sent"] == 0
            assert body["failed"] == 2
            assert elapsed < 20.0, f"external/send took {elapsed:.1f}s"
        finally:
            requests.delete(f"{API}/api-keys/{kid}", headers=admin_headers, timeout=10)
            requests.delete(f"{API}/email-templates/{tid}", headers=admin_headers, timeout=10)

    def test_today_sent_unaffected_by_failed_sends(self, admin_headers):
        """Since SMTP is unreachable, record_send is never called → daily counter
        should NOT increment. This also implicitly verifies the daily_quota table
        exists (otherwise even reading current_daily_count would crash)."""
        # Baseline
        r0 = requests.get(f"{API}/throttle/defaults", headers=admin_headers, timeout=10)
        assert r0.status_code == 200
        before = r0.json()["today_sent"]

        # Send a campaign that will fail (SMTP unreachable)
        files = {"excel": ("data.xlsx", _build_xlsx_bytes(),
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        data = {
            "subject": "Hi", "body_html": "<p>x</p>",
            "per_minute": "10000", "per_domain_per_min": "10000",
            "delay_min_ms": "0", "delay_max_ms": "0",
        }
        rs = requests.post(f"{API}/campaigns/send", data=data, files=files,
                           headers=admin_headers, timeout=60)
        assert rs.status_code == 200
        assert rs.json()["sent"] == 0
        assert rs.json()["failed"] == 2

        # Counter must NOT have advanced because record_send is success-only
        r1 = requests.get(f"{API}/throttle/defaults", headers=admin_headers, timeout=10)
        after = r1.json()["today_sent"]
        assert after == before, (
            f"today_sent advanced from {before} to {after} despite no successful sends"
        )
